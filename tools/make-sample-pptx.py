"""生成 test\\纯页码课件.pptx —— 10 页的纯页码示例课件。

这份课件是给「课件续页」插件当试验品用的：每页只有一个巨大的页码和一个深色背景，
页面上不写任何说明文字，于是插件显示"第 N 页"和幻灯片实际内容永远对得上。

重新生成（仓库根目录，用带 python-pptx 的解释器）：
    python tools\\make-sample-pptx.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SLIDE_COUNT = 10

# 十种偏深的纯色，白字在上面足够清楚；相邻两页色相差别大，翻页时一眼能看出变了。
BACKGROUNDS = [
    "1F3A5F",  # 深蓝
    "3A1F5F",  # 深紫
    "5F1F2E",  # 酒红
    "1F5F4A",  # 深青绿
    "5F4A1F",  # 深褐
    "1F2E5F",  # 靛蓝
    "4A1F5F",  # 紫罗兰
    "1F5F2E",  # 深绿
    "5F2E1F",  # 砖红
    "2E1F1F",  # 近黑
]

# 数字的字号：填满大半个页面，投影到教室后排也看得清。
NUMBER_PT = 400


def build(out_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]  # 只有版式、没有任何占位符

    for index in range(SLIDE_COUNT):
        slide = presentation.slides.add_slide(blank)

        # 背景：整页纯色
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(BACKGROUNDS[index % len(BACKGROUNDS)])

        # 正文：一个铺满全页的文本框，只有一个页码，居中
        box = slide.shapes.add_textbox(Emu(0), Emu(0), presentation.slide_width, presentation.slide_height)
        frame = box.text_frame
        # word_wrap=True + auto_size=NONE 是有意的：python-pptx 默认给文本框写
        # wrap="none" + spAutoFit，那会让放映时文本框缩到文字大小、从左上角开始摆，
        # 数字就不在页面正中间了。这里要的是"铺满整页、文字居中、尺寸固定"。
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE

        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = str(index + 1)
        run.font.size = Pt(NUMBER_PT)
        run.font.bold = True
        run.font.name = "Microsoft YaHei"
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(out_path)


if __name__ == "__main__":
    target = Path(__file__).resolve().parent.parent / "test" / "纯页码课件.pptx"
    build(target)
    print(f"已生成 {target}")
